"""Exports each module deck as .pptx, ready to open in Google Slides.

    npm run build:slides:pptx
    python3 scripts/slides/build_pptx.py clustering

Google Slides cannot import a PDF as slides -- only .pptx -- so this renders the
same template with native shapes and text boxes. Everything stays editable
after import: the headings, body copy, code and bullets are real text in Roboto
and Roboto Mono, both of which Google Slides has built in.

Display equations are placed as images (there is no equation object Slides
understands); inline maths is written as Unicode text so the sentence around it
stays editable.
"""
import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Pt

sys.path.insert(0, str(Path(__file__).parent))
import build_slides as B  # noqa: E402
import template as T  # noqa: E402

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "slides"


def rgb(triple):
    return RGBColor(*[int(round(c * 255)) for c in triple])


PURPLE, BLUE, PANEL_BLUE = rgb(T.PURPLE), rgb(T.BLUE), rgb(T.PANEL_BLUE)
TITLE_INK, BODY_INK, WHITE = rgb(T.TITLE_INK), rgb(T.BODY_INK), rgb(T.WHITE)
SUBTITLE, CODE_BG, CODE_EDGE = rgb(T.SUBTITLE), rgb(T.CODE_BG), rgb(T.CODE_EDGE)

FONT_SANS, FONT_MONO = "Roboto", "Roboto Mono"


def _no_line(shape):
    shape.line.fill.background()
    shape.shadow.inherit = False


def rect(slide, x, y, w, h, fill):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Pt(x), Pt(y), Pt(w), Pt(h))
    s.fill.solid(); s.fill.fore_color.rgb = fill
    _no_line(s)
    return s


def dot(slide, x, y, d, fill):
    s = slide.shapes.add_shape(MSO_SHAPE.OVAL, Pt(x), Pt(y), Pt(d), Pt(d))
    s.fill.solid(); s.fill.fore_color.rgb = fill
    _no_line(s)
    return s


def textbox(slide, x, y, w, h, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.TOP
    return tf


def run(para, text, size, colour, font=FONT_SANS, bold=False):
    r = para.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.name = font
    r.font.bold = bold
    r.font.color.rgb = colour
    return r


def dot_triangle(slide, x0, y0, d, pitch, rows, colour):
    for r in range(rows):
        for c in range(rows - r):
            dot(slide, x0 + (c + r) * pitch, y0 + r * pitch, d, colour)


def logo(slide, x, y, size):
    """The purple ring with a three-node graph, drawn as native shapes."""
    ring = slide.shapes.add_shape(MSO_SHAPE.OVAL, Pt(x), Pt(y), Pt(size), Pt(size))
    ring.fill.solid(); ring.fill.fore_color.rgb = WHITE
    ring.line.color.rgb = PURPLE
    ring.line.width = Pt(max(1.2, size * 0.06))
    ring.shadow.inherit = False
    r = size / 2
    nodes = [(x + r, y + r * 0.58), (x + r * 0.58, y + r * 1.40), (x + r * 1.42, y + r * 1.40)]
    for nx, ny in nodes:
        n = size * 0.16
        d = slide.shapes.add_shape(MSO_SHAPE.OVAL, Pt(nx - n / 2), Pt(ny - n / 2), Pt(n), Pt(n))
        d.fill.solid(); d.fill.fore_color.rgb = PURPLE
        _no_line(d)


def footer(slide):
    tf = textbox(slide, 560, 366, 118, 14, PP_ALIGN.RIGHT)
    run(tf.paragraphs[0], "CP Initiative", 11, BODY_INK, bold=True)
    tf2 = textbox(slide, 560, 378, 118, 12, PP_ALIGN.RIGHT)
    run(tf2.paragraphs[0], "joincpi.org", 8, BODY_INK)
    logo(slide, 683.5, 368.0, 27.4)


def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def content_slide(prs, title):
    s = blank(prs)
    rect(s, 0, 0, T.W, T.TOPBAR_H, PURPLE)
    dot_triangle(s, 646.5, 23.1, 10.0, 18.0, 4, BLUE)
    if title:
        tf = textbox(s, T.X_MARGIN, T.TITLE_Y - 6, T.W - 2 * T.X_MARGIN - 90, 40)
        run(tf.paragraphs[0], title, T.T_TITLE, TITLE_INK, bold=True)
    footer(s)
    return s


def title_slide(prs, title, subtitle):
    s = blank(prs)
    rect(s, 0, 0, 447.1, T.H, PANEL_BLUE)
    lines = T._wrap(title, "bold", T.T_TITLE_HERO, 470.0)
    band_h = 162.2 if len(lines) <= 2 else 162.2 + (len(lines) - 2) * 40.0
    band_y = 109.4 - (band_h - 162.2) / 2
    rect(s, 0, band_y, 554.7, band_h, PURPLE)
    for i in range(5):
        rect(s, 0, 144.5 + i * 21.15, 25.1, 9.2, WHITE)
    dot_triangle(s, 18.4, 20.5, 10.0, 18.0, 4, BLUE)
    for r in range(8):
        for c in range(5):
            dot(s, 565.0 + c * 13.1, 109.4 + r * 22.44, 5.1, PURPLE)

    tf = textbox(s, 61.0, band_y + 22, 470, band_h - 60)
    p = tf.paragraphs[0]
    run(p, title, T.T_TITLE_HERO, WHITE, bold=True)
    sub = textbox(s, 61.0, band_y + band_h - 46, 470, 30)
    size = T.T_SUB_HERO
    while size > 11 and T.width(subtitle, "regular", size) > 470.0:
        size -= 0.5
    run(sub.paragraphs[0], subtitle, size, SUBTITLE)

    logo(s, 57.3, 312.4, 47.0)
    tf = textbox(s, 109.7, 316, 200, 24)
    run(tf.paragraphs[0], "CP Initiative", 18, BODY_INK, bold=True)
    tf = textbox(s, 112.3, 340, 200, 18)
    run(tf.paragraphs[0], "joincpi.org", 12, BODY_INK)
    return s


class PptxDeck:
    """Same pagination as the PDF builder, emitting pptx shapes."""

    def __init__(self, prs, module):
        self.prs, self.module = prs, module
        self.slide, self.y, self.heading = None, 0.0, ""

    def start(self, heading):
        self.heading = heading
        self.slide = content_slide(self.prs, heading)
        self.y = T.BODY_TOP if heading else T.TITLE_Y

    def room(self, need):
        return self.y + need <= T.BODY_MAX_Y

    def ensure(self, need):
        if self.slide is None or not self.room(need):
            self.start(self.heading)

    def _plain(self, text):
        """Flatten inline runs to text; maths becomes Unicode so it stays editable."""
        parts = []
        for chunk, style in B.inline_runs(text):
            parts.append(B.tex_to_text(chunk) if style == "math" else chunk)
        return "".join(parts)

    def paragraph(self, text):
        body = self._plain(text)
        lines = len(B.wrap_runs([(body, "regular")], T.T_BODY, T.W - 2 * T.X_MARGIN))
        need = lines * T.LEAD_BODY
        self.ensure(min(need, 3 * T.LEAD_BODY))
        h = min(need, T.BODY_MAX_Y - self.y)
        tf = textbox(self.slide, T.X_MARGIN, self.y, T.W - 2 * T.X_MARGIN, h)
        tf.paragraphs[0].line_spacing = T.LEAD_BODY / T.T_BODY
        run(tf.paragraphs[0], body, T.T_BODY, BODY_INK)
        self.y += need + 10

    def bullets(self, items, ordered=False):
        for n, item in enumerate(items, 1):
            body = self._plain(item)
            lines = len(B.wrap_runs([(body, "regular")], T.T_BODY,
                                    T.W - B.BULLET_TEXT_X - T.X_MARGIN))
            need = lines * T.LEAD_BODY
            self.ensure(min(need, 2 * T.LEAD_BODY))
            marker = textbox(self.slide, B.BULLET_X, self.y, 30, T.LEAD_BODY)
            run(marker.paragraphs[0], f"{n}." if ordered else "●",
                T.T_BODY if ordered else 11, BODY_INK)
            tf = textbox(self.slide, B.BULLET_TEXT_X, self.y,
                         T.W - B.BULLET_TEXT_X - T.X_MARGIN, need)
            tf.paragraphs[0].line_spacing = T.LEAD_BODY / T.T_BODY
            run(tf.paragraphs[0], body, T.T_BODY, BODY_INK)
            self.y += need + 6
        self.y += 6

    def code(self, lines):
        if not lines:
            return
        h = len(lines) * T.LEAD_CODE + 22
        self.ensure(h)
        if not self.room(h):
            self.start(self.heading)
        panel = rect(self.slide, T.X_MARGIN - 6, self.y,
                     T.W - 2 * T.X_MARGIN + 12, h, CODE_BG)
        panel.line.color.rgb = CODE_EDGE
        panel.line.width = Pt(0.8)
        tf = textbox(self.slide, T.X_MARGIN, self.y + 10,
                     T.W - 2 * T.X_MARGIN, h - 16)
        for i, line in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.line_spacing = T.LEAD_CODE / T.T_CODE
            for tok, col in B.code_runs(line):
                run(p, tok, T.T_CODE, rgb(col), font=FONT_MONO)
        self.y += h + 12

    def display_math(self, tex):
        img = B.math_image(tex, T.T_BODY + 2)
        if img is None:
            self.paragraph(B.tex_to_text(tex))
            return
        data, w, h = img
        if w > T.W - 2 * T.X_MARGIN:
            h *= (T.W - 2 * T.X_MARGIN) / w
            w = T.W - 2 * T.X_MARGIN
        self.ensure(h + 18)
        if not self.room(h + 18):
            self.start(self.heading)
        import io
        self.slide.shapes.add_picture(io.BytesIO(data), Pt((T.W - w) / 2),
                                      Pt(self.y + 6), Pt(w), Pt(h))
        self.y += h + 18

    def table(self, tbl):
        head, rows = tbl.get("head", []), tbl.get("rows", [])
        if not head:
            return
        h = (len(rows) + 1) * 24 + 8
        self.ensure(h)
        if not self.room(h):
            self.start(self.heading)
        shape = self.slide.shapes.add_table(len(rows) + 1, len(head),
                                            Pt(T.X_MARGIN), Pt(self.y),
                                            Pt(T.W - 2 * T.X_MARGIN), Pt(h))
        tbl_obj = shape.table
        for c, cell in enumerate(head):
            tc = tbl_obj.cell(0, c).text_frame
            run(tc.paragraphs[0], self._plain(cell), 12, TITLE_INK, bold=True)
        for r, row in enumerate(rows, 1):
            for c, cell in enumerate(row[:len(head)]):
                tc = tbl_obj.cell(r, c).text_frame
                run(tc.paragraphs[0], self._plain(cell), 12, BODY_INK)
        self.y += h + 12

    def image(self, src):
        path = ROOT / src
        if not path.exists():
            return
        import io
        from PIL import Image
        im = Image.open(path)
        avail_w = T.W - 2 * T.X_MARGIN
        avail_h = T.BODY_MAX_Y - T.BODY_TOP - 10
        scale = min(avail_w / im.width, avail_h / im.height)
        w, h = im.width * scale, im.height * scale
        self.start(self.heading)
        buf = io.BytesIO()
        im.convert("RGB").resize((int(w * 2), int(h * 2)), Image.LANCZOS).save(
            buf, "PNG", optimize=True)
        buf.seek(0)
        self.slide.shapes.add_picture(buf, Pt((T.W - w) / 2), Pt(self.y), Pt(w), Pt(h))
        self.y += h + 10


def build(module, lesson):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(Pt(T.W)), Emu(Pt(T.H))
    title_slide(prs, module["title"], f"{B.SITE}/module/?id={module['id']}")
    deck = PptxDeck(prs, module)
    for section in lesson["sections"]:
        if section["id"] == "prerequisites":
            continue
        if len(prs.slides._sldIdLst) >= B.SLIDE_BUDGET:
            break
        deck.start(section["title"] or module["title"])
        for block in B.condense(section):
            if len(prs.slides._sldIdLst) >= B.SLIDE_BUDGET:
                break
            if "p" in block:
                deck.paragraph(block["p"])
            elif "heading" in block:
                deck.start(block["heading"])
            elif "list" in block:
                deck.bullets(block["list"])
            elif "olist" in block:
                deck.bullets(block["olist"], ordered=True)
            elif "code" in block:
                deck.code(block["code"].get("lines", []))
            elif "tex" in block:
                deck.display_math(block["tex"])
            elif "table" in block:
                deck.table(block["table"])
            elif "image" in block:
                deck.image(block["image"]["src"])
            elif "quote" in block:
                deck.paragraph(" ".join(block["quote"]))
    OUT.mkdir(exist_ok=True)
    prs.save(str(OUT / f"{module['id']}.pptx"))
    return len(prs.slides._sldIdLst)


def main():
    data = B.load_content()
    only = sys.argv[1] if len(sys.argv) > 1 else None
    total = 0
    for module in data["modules"]:
        if only and only != module["id"]:
            continue
        lesson = data["lessons"].get(module["id"])
        if lesson:
            n = build(module, lesson)
            total += n
            print(f"  {n:3d} slides  {module['id']}.pptx")
    print(f"\n{total} slides written to slides/")


if __name__ == "__main__":
    main()
